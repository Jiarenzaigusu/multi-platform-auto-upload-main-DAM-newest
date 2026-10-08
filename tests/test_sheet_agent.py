from pathlib import Path
from datetime import date
from unittest.mock import MagicMock

from openpyxl import Workbook, load_workbook

from webapp.sheet_agent import (
    apply_plan,
    column_group_labels,
    direct_header_correction_plan,
    direct_coordinate_correction_plan,
    fallback_append_plan,
    find_header,
    matched_source_headers,
    plan_field_options,
    profile_field_options,
    remove_target_mappings,
    source_profile,
    template_profile,
    unmapped_coordinate_targets,
)
from webapp.sheet_database_source import export_content_source


def _workbook(path: Path, rows: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def test_database_source_exports_only_selected_brand_and_period(tmp_path):
    database = MagicMock()
    database.execute.side_effect = [
        [("brand_id",), ("年份",), ("下载周期",), ("内容ID",), ("点击次数",)],
        [(2026, "8月", "1383899085827225", 53)],
    ]
    path = export_content_source(database, 7, "星巴克", "2026-08", tmp_path / "数据库内容表现.xlsx")

    sheet = load_workbook(path).active
    assert list(sheet.values) == [
        ("品牌", "年份", "下载周期", "内容ID", "点击次数"),
        ("星巴克", 2026, "8月", "1383899085827225", 53),
    ]
    assert source_profile(path)["sheets"][0]["headers"] == ["品牌", "年份", "下载周期", "内容ID", "点击次数"]
    query, parameters = database.execute.call_args_list[1].args
    assert "WHERE brand_id = %s AND `年份` = %s AND `下载周期` = %s" in query
    assert parameters == (7, 2026, "8月")


def test_append_rows_fills_matching_columns_without_overwriting_formula(tmp_path):
    source = tmp_path / "source.xlsx"
    template = tmp_path / "template.xlsx"
    output = tmp_path / "output.xlsx"
    _workbook(source, [["商品", "金额"], ["Polo", 128]])
    _workbook(template, [["商品名称", "成交金额", "保护公式"], [None, None, "=1+1"]])
    plan = {
        "rules": [{
            "mode": "append_rows",
            "source_file": source.name,
            "source_sheet": "Sheet",
            "target_sheet": "Sheet",
            "source_header_row": 1,
            "target_header_row": 1,
            "column_map": [
                {"source_header": "商品", "target_header": "商品名称"},
                {"source_header": "金额", "target_header": "成交金额"},
            ],
            "confidence": 0.99,
            "reason": "字段语义一致",
        }]
    }

    audit = apply_plan(template, [source], plan, output, 0.9)

    sheet = load_workbook(output, data_only=False).active
    assert (sheet["A3"].value, sheet["B3"].value) == ("Polo", 128)
    assert sheet["C2"].value == "=1+1"
    assert audit[0]["status"] == "applied"


def test_fallback_append_plan_maps_safe_header_aliases(tmp_path):
    source = tmp_path / "source.xlsx"
    template = tmp_path / "template.xlsx"
    _workbook(source, [["内容ID", "内容名称", "内容发布时间", "查看人数"], ["1", "Polo", "2026-01-01", 8]])
    _workbook(template, [["内容id", "标签", "名称", "发布时间"]])

    plan = fallback_append_plan({
        "sources": [source_profile(source)],
        "template": template_profile(template),
    })

    assert plan["rules"][0]["column_map"] == [
        {"source_header": "内容ID", "source_header_cell": "A1", "target_header": "内容id", "target_header_cell": "A1"},
        {"source_header": "内容名称", "source_header_cell": "B1", "target_header": "名称", "target_header_cell": "C1"},
        {"source_header": "内容发布时间", "source_header_cell": "C1", "target_header": "发布时间", "target_header_cell": "D1"},
    ]


def test_header_corrections_replace_selected_mapping_and_keep_others(tmp_path):
    source = tmp_path / "source.xlsx"
    template = tmp_path / "template.xlsx"
    _workbook(source, [["原标题", "发布时间"], ["Polo", "2026-01-01"]])
    _workbook(template, [["名称", "日期"]])
    profile = {"sources": [source_profile(source)], "template": template_profile(template)}
    original = {"rules": [{
        "mode": "append_rows", "column_map": [
            {"source_header": "原标题", "target_header": "名称"},
            {"source_header": "发布时间", "target_header": "日期"},
        ],
    }]}

    remaining = remove_target_mappings(original, ["名称"])
    correction = direct_header_correction_plan(profile, ["原标题"], ["名称"])

    assert matched_source_headers(remaining) == ["发布时间"]
    assert correction["rules"][0]["column_map"] == [
        {"source_header": "原标题", "target_header": "名称"}
    ]


def test_coordinate_mapping_distinguishes_duplicate_headers(tmp_path):
    source = tmp_path / "source.xlsx"
    template = tmp_path / "template.xlsx"
    output = tmp_path / "output.xlsx"
    _workbook(source, [["金额", "金额"], [1, 2]])
    _workbook(template, [["金额", "金额"]])
    plan = direct_coordinate_correction_plan([{
        "source": {"file": source.name, "sheet": "Sheet", "header_row": 1, "cell": "B1", "header": "金额"},
        "target": {"sheet": "Sheet", "header_row": 1, "cell": "A1", "header": "金额"},
    }])

    apply_plan(template, [source], plan, output, 0.9)

    assert load_workbook(output).active["A2"].value == 2


def test_unmapped_coordinate_targets_keeps_only_targets_missing_from_plan():
    plan = {"rules": [
        {
            "mode": "append_rows", "target_sheet": "明细",
            "column_map": [{"target_header_cell": "B3"}],
        },
        {
            "mode": "metric_fill", "target_sheet": "周报",
            "target_label_cell": "A8",
        },
    ]}
    targets = [
        {"sheet": "明细", "cell": "B3", "header": "名称"},
        {"sheet": "明细", "cell": "C3", "header": "金额"},
        {"sheet": "周报", "cell": "A8", "header": "曝光次数"},
        {"sheet": "周报", "cell": "A9", "header": "播放人数"},
    ]

    assert unmapped_coordinate_targets(plan, targets) == [
        {"sheet": "明细", "cell": "C3", "header": "金额"},
        {"sheet": "周报", "cell": "A9", "header": "播放人数"},
    ]


def test_find_header_accepts_model_coordinate_prefix(tmp_path):
    workbook_path = tmp_path / "headers.xlsx"
    _workbook(workbook_path, [[None], [None], [None], [None], [None], ["统计日期", "成交金额"]])
    sheet = load_workbook(workbook_path).active

    assert find_header(sheet, "A6:统计日期", 6) == 1
    assert find_header(sheet, "B6:成交金额", 6) == 2


def test_saved_metric_plan_recovers_coordinate_field_options():
    options = plan_field_options({"rules": [{
        "mode": "metric_fill", "source_file": "source.xls", "source_sheet": "数据",
        "source_date_header": "A6:日期", "source_value_header": "C6:成交金额",
        "target_sheet": "周报", "target_label_cell": "D12", "target_label": "种草金额",
        "target_row": 12,
    }]})

    assert [item["cell"] for item in options["sources"]] == ["A6", "C6"]
    assert options["targets"][0]["label"] == "周报 / D12 / 种草金额"


def test_template_fields_carry_the_block_above_so_repeated_names_stay_distinct(tmp_path):
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "周报"
    sheet.merge_cells("D1:L11")
    sheet["D1"] = "订阅：本周新增商家发布订阅数量21条，曝光人数104,127，环比上涨5.39%"
    sheet.merge_cells("M1:U11")
    sheet["M1"] = "内容端上周整体成交金额524,000，占比全店支付金额的7.57%"
    sheet["K12"] = "TTL"
    sheet["L12"] = "环比"
    sheet["T12"] = "TTL"
    sheet["U12"] = "环比"
    workbook.save(template)

    sheet_model = template_profile(template)["sheets"][0]
    groups = next(item for item in sheet_model["table_header_candidates"] if item["header_row"] == 12)["column_groups"]

    assert groups[11] == "订阅 D1:L11"
    assert groups[20] == "上方 M1:U11"

    options = profile_field_options({"sources": [], "template": {"file": "template.xlsx", "sheets": [sheet_model]}})
    labels = {item["cell"]: item["label"] for item in options["targets"]}

    assert labels["K12"] == "周报 / K12 / 订阅 D1:L11 · TTL"
    assert labels["T12"] == "周报 / T12 / 上方 M1:U11 · TTL"
    assert len(set(labels.values())) == len(labels)
    assert [item["header"] for item in options["targets"] if item["cell"] == "K12"] == ["TTL"]


def test_header_columns_without_a_block_above_get_no_group(tmp_path):
    template = tmp_path / "flat.xlsx"
    _workbook(template, [["内容ID", "内容名称", "曝光人数"]])

    sheet_model = template_profile(template)["sheets"][0]
    options = profile_field_options({"sources": [], "template": {"file": "flat.xlsx", "sheets": [sheet_model]}})

    assert [item["label"] for item in options["targets"]] == ["Sheet / A1 / 内容ID", "Sheet / B1 / 内容名称", "Sheet / C1 / 曝光人数"]
    assert all(item["group"] is None for item in options["targets"])
    assert column_group_labels(load_workbook(template).active, None, 1, 3) == {}


def test_metric_targets_use_visible_sheet_row_labels_instead_of_hidden_headers(tmp_path):
    template = tmp_path / "report.xlsx"
    workbook = Workbook()
    old = workbook.active
    old.title = "历史内容数据"
    old.sheet_state = "hidden"
    old.append(["渠道", "数据/规划", "日期", "TTL", "环比"])
    report = workbook.create_sheet("内容报表")
    report.append(["渠道", "数据/规划", "日期", date(2026, 9, 1), date(2026, 9, 2), date(2026, 9, 3), "TTL"])
    report.append(["内容端", "实际数据", "短视频种草金额", 10, 20, 30, "=SUM(D2:F2)"])
    report.append([None, None, "内容曝光人数", 100, 200, 300, "=SUM(D3:F3)"])
    workbook.save(template)

    options = profile_field_options({"sources": [], "template": template_profile(template)})

    assert [(item["sheet"], item["cell"], item["header"]) for item in options["targets"]] == [
        ("内容报表", "C2", "短视频种草金额"),
        ("内容报表", "C3", "内容曝光人数"),
    ]
