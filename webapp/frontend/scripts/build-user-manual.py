"""Export the saved Word manual to self-contained, locally hosted web help.

Run with the Codex bundled Python runtime (python-docx is required).
No network calls or changes to the input document are made.
"""

import argparse
import html
import re
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


FRONTEND = Path(__file__).resolve().parents[1]
PROJECT = FRONTEND.parents[1]
DEFAULT_SOURCE = PROJECT / "outputs/manual/智能发布中枢系统用户使用手册.docx"


def build(source: Path, output: Path) -> None:
    document = Document(source)
    output.mkdir(parents=True, exist_ok=True)
    image_dir = output / "images"
    image_dir.mkdir(exist_ok=True)
    images = {}
    image_sizes = {}
    headings = []
    sections = []
    current = []
    current_id = "overview"
    current_label = "手册概览"
    bookmark_targets = {}
    for paragraph in document.paragraphs:
        if paragraph.style.name == "Heading 1":
            target = f"section-{len(headings) + 1}"
            headings.append((target, paragraph.text))
            for bookmark in paragraph._p.xpath(".//w:bookmarkStart"):
                bookmark_targets[bookmark.get(qn("w:name"))] = target

    def run_content(element):
        parts = []
        for child in element:
            if child.tag == qn("w:t"):
                parts.append(html.escape(child.text or ""))
            elif child.tag == qn("w:br"):
                if child.get(qn("w:type")) != "page":
                    parts.append("<br>")
            elif child.tag == qn("w:tab"):
                parts.append(" ")
            elif child.tag == qn("w:drawing"):
                descriptions = child.xpath(".//wp:docPr")
                description = descriptions[0].get("descr", "操作界面截图") if descriptions else "操作界面截图"
                for blip in child.xpath(".//a:blip"):
                    relationship = blip.get(qn("r:embed"))
                    if not relationship:
                        continue
                    if relationship not in images:
                        part = document.part.related_parts[relationship]
                        name = f"screenshot-{len(images) + 1}{Path(str(part.partname)).suffix}"
                        (image_dir / name).write_bytes(part.blob)
                        images[relationship] = name
                        image_sizes[relationship] = (part.image.px_width, part.image.px_height)
                    path = f"images/{images[relationship]}"
                    width, height = image_sizes[relationship]
                    ratio = width / height
                    shape = "portrait" if ratio < 0.8 else "balanced" if ratio < 1.4 else "wide"
                    parts.append(f'<a class="screenshot-link screenshot-{shape}" style="--image-ratio:{ratio:.6f}" href="{path}" target="_blank" rel="noopener" aria-label="放大查看截图" aria-haspopup="dialog"><img src="{path}" width="{width}" height="{height}" alt="{html.escape(description, quote=True)}" loading="lazy"></a>')
        text = "".join(parts)
        properties = element.find(qn("w:rPr"))
        if properties is not None and text and "<img " not in text:
            for tag, property_name in (("strong", "b"), ("em", "i")):
                prop = properties.find(qn(f"w:{property_name}"))
                if prop is not None and prop.get(qn("w:val"), "1") not in ("0", "false"):
                    text = f"<{tag}>{text}</{tag}>"
        return text

    def paragraph_content(paragraph):
        parts = []
        for child in paragraph._p:
            if child.tag == qn("w:r"):
                parts.append(run_content(child))
            elif child.tag == qn("w:hyperlink"):
                content = "".join(run_content(run) for run in child.findall(qn("w:r")))
                anchor = child.get(qn("w:anchor"))
                if anchor:
                    target = bookmark_targets.get(anchor, "overview")
                    content = re.sub(r"\s*第\s*\d+\s*页", "", content)
                    parts.append(f'<a href="#{target}">{content}</a>')
                else:
                    relationship = child.get(qn("r:id"))
                    url = document.part.rels[relationship].target_ref if relationship else ""
                    if str(url).startswith(("https://", "http://", "mailto:")):
                        parts.append(f'<a href="{html.escape(str(url), quote=True)}" target="_blank" rel="noopener">{content}</a>')
                    else:
                        parts.append(content)
        return "".join(parts)

    def render_paragraph(paragraph):
        if paragraph.style.name == "Title":
            return ""
        content = paragraph_content(paragraph)
        if not content.strip():
            return ""
        if "Word 的" in paragraph.text:
            content = "点击上表或章节目录，即可跳转到对应说明。点击截图可放大、缩小，或按原始尺寸查看。"
        style = paragraph.style.name
        if style == "Command":
            return f'<p class="command"><code>{content}</code></p>'
        if style == "Heading 2":
            return f"<h3>{content}</h3>"
        if style == "Caption":
            return f'<p class="caption">{content}</p>'
        if style == "List Bullet":
            return f'<p class="bullet">{content}</p>'
        if style == "Subtitle":
            return f'<p class="subtitle">{content}</p>'
        class_attribute = ' class="figure"' if "<img " in content else ""
        return f"<p{class_attribute}>{content}</p>"

    def render_table(element):
        rows = []
        for index, row in enumerate(Table(element, document).rows):
            tag = "th" if index == 0 else "td"
            cells = []
            for cell in row.cells:
                scope = ' scope="col"' if tag == "th" else ""
                cells.append(f"<{tag}{scope}>" + "".join(render_paragraph(p) for p in cell.paragraphs) + f"</{tag}>")
            rows.append("<tr>" + "".join(cells) + "</tr>")
        return '<div class="table-scroll" tabindex="0" role="region" aria-label="可横向滚动的说明表"><table><thead>' + rows[0] + "</thead><tbody>" + "".join(rows[1:]) + "</tbody></table></div>"

    heading_index = 0
    for element in document.element.body:
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, document)
            if paragraph.style.name == "Heading 1":
                sections.append((current_id, current_label, "".join(current)))
                current_id, current_label = headings[heading_index]
                heading_index += 1
                current = [f"<h2>{html.escape(current_label)}</h2>"]
            else:
                current.append(render_paragraph(paragraph))
        elif element.tag == qn("w:tbl"):
            current.append(render_table(element))
    sections.append((current_id, current_label, "".join(current)))
    chapter_titles = {
        "1": "首次使用", "2": "平台账号登录", "3": "单条发布与素材",
        "4": "平台填写规则", "5": "批量发布", "6": "AI 文案",
        "7": "任务与结果", "8": "常见问题", "9": "日常维护",
        "A": "部署与管理", "B": "数据安全与恢复", "C": "名词与项目位置",
    }
    groups = {}
    for target, label, _ in sections:
        match = re.match(r"^(\d+)\s+(.+)$", label) or re.match(r"^附录\s+([A-Z])\s+(.+)$", label)
        key, child_label = (match.group(1), match.group(2)) if match else (target, label)
        if key == "1":
            child_label = re.sub(r"^首次使用\s+", "", child_label)
        groups.setdefault(key, []).append((target, child_label))
    toc_parts = []
    for key, children in groups.items():
        label = chapter_titles.get(key, children[0][1])
        number = key.zfill(2) if key.isdigit() else key if key in ("A", "B", "C") else ""
        badge = f'<span class="toc-number" aria-hidden="true">{number}</span>' if number else ""
        heading = badge + f'<span class="toc-label">{html.escape(label)}</span>'
        if key == "A":
            toc_parts.append('<p class="toc-divider">管理员参考</p>')
        if len(children) == 1:
            target = children[0][0]
            toc_parts.append(f'<a class="toc-chapter-link" href="#{target}" data-chapter="{key}" data-label="{html.escape(label, quote=True)}">{heading}</a>')
        else:
            child_links = "".join(f'<a class="toc-child-link" href="#{target}" data-label="{html.escape(child, quote=True)}">{html.escape(child)}</a>' for target, child in children)
            toc_parts.append(f'<details class="toc-group" data-chapter="{key}" data-label="{html.escape(label, quote=True)}"><summary>{heading}</summary><div class="toc-children">{child_links}</div></details>')
    toc = "".join(toc_parts)
    body = "".join(f'<section id="{target}" class="manual-section">{content}</section>' for target, _, content in sections)
    page = f'''<!doctype html>
<html lang="zh-CN">
<head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>使用手册 · 智能发布中枢系统</title><link rel="stylesheet" href="manual.css"><script src="manual.js" defer></script></head>
<body><a class="skip-link" href="#manual-content">跳到手册正文</a>
<header class="manual-header"><div><p class="brand">智能发布中枢系统</p><h1>使用手册</h1></div>
<div class="header-actions"><a href="/">返回发布台</a><a class="download" href="user-manual.docx" download="智能发布中枢系统用户使用手册.docx">下载 Word 版</a></div></header>
<div class="manual-layout"><aside class="manual-sidebar"><label for="manual-search">查找操作说明</label>
<input id="manual-search" type="search" placeholder="例如：配对、批量、失败" autocomplete="off">
<p id="search-status" role="status" aria-live="polite">输入关键词查找相关操作。</p>
<details id="manual-toc" open><summary>章节目录</summary>
<div class="toc-controls"><button type="button" data-toc-action="expand">全部展开</button><button type="button" data-toc-action="collapse">全部收起</button></div>
<nav aria-label="手册章节">{toc}</nav></details></aside>
<main id="manual-content" tabindex="-1">{body}<p id="no-results" hidden>没有找到相关章节，请尝试其他关键词。</p></main></div>
<dialog id="image-viewer" aria-labelledby="image-viewer-title">
<header class="image-viewer-toolbar"><h2 id="image-viewer-title">查看截图</h2>
<div class="image-viewer-actions"><button type="button" data-image-action="out" aria-label="缩小截图">−</button>
<output id="image-zoom" aria-live="polite"></output><button type="button" data-image-action="in" aria-label="放大截图">＋</button>
<button type="button" data-image-action="fit">适应窗口</button><button type="button" data-image-action="original">原始尺寸</button>
<button type="button" data-image-action="close" autofocus>关闭</button></div></header>
<div class="image-viewer-viewport" tabindex="0" aria-label="截图查看区域 放大后可滚动查看"><div class="image-viewer-stage"><img id="image-viewer-image" alt=""></div></div>
</dialog>
<footer>在线内容与本页 Word 下载来自同一份手册。平台界面如有变化，以实际页面与任务日志为准。</footer></body></html>'''
    (output / "index.html").write_text(page, encoding="utf-8")
    shutil.copy2(source, output / "user-manual.docx")
    print(f"Exported {len(sections)} sections, {len(images)} screenshots to {output}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, nargs="?", default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=FRONTEND / "public/assets/user-manual")
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve())
